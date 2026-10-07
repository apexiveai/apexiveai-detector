import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pymupdf
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from PIL import Image, ImageDraw

from app.engine import (
    cv,
    pdf_assets,
    prepare_verification,
    run,
    verify,
    verify_prepared,
    xlsx_assets,
)


class MarkAssetExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        image = Image.new('RGB', (40, 30), 'white')
        drawing = ImageDraw.Draw(image)
        drawing.ellipse((3, 2, 22, 28), fill='black')
        drawing.rectangle((21, 9, 37, 21), fill='black')
        self.image_path = self.root / 'logo.png'
        image.save(self.image_path)
        self.image_bytes = self.image_path.read_bytes()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_pdf_extracts_only_images_in_the_540_field(self):
        pdf_path = self.root / 'sample.pdf'
        doc = pymupdf.open()
        page = doc.new_page(width=300, height=300)
        page.insert_text((20, 24), 'Application No. T/2026/001111')
        page.insert_text((20, 40), '(540)')
        page.insert_image(pymupdf.Rect(120, 50, 200, 100), stream=self.image_bytes)
        page.insert_text((20, 140), '(550)')
        page.insert_image(pymupdf.Rect(120, 180, 200, 230), stream=self.image_bytes)
        doc.save(pdf_path)
        doc.close()

        assets = pdf_assets(pdf_path, self.root, lambda *_: None)

        self.assertEqual(len(assets), 1)
        self.assertEqual(assets[0]['field_code'], '(540)')
        self.assertEqual(assets[0]['extract_method'], 'pdf_540_embedded_image')
        self.assertEqual(assets[0]['application_no'], 'T/2026/001111')
        self.assertTrue((self.root / assets[0]['image_file']).is_file())

    def test_pdf_logo_fields_keep_their_own_application_numbers(self):
        pdf_path = self.root / 'multiple_records.pdf'
        doc = pymupdf.open()
        page = doc.new_page(width=300, height=300)
        page.insert_text((20, 24), 'Application No. T/2026/001111')
        page.insert_text((20, 40), '(540)')
        page.insert_image(pymupdf.Rect(120, 50, 200, 100), stream=self.image_bytes)
        page.insert_text((20, 124), 'Application No. T/2026/002222')
        page.insert_text((20, 140), '(540)')
        page.insert_image(pymupdf.Rect(120, 170, 200, 220), stream=self.image_bytes)
        doc.save(pdf_path)
        doc.close()

        assets = pdf_assets(pdf_path, self.root, lambda *_: None)

        self.assertEqual([asset['application_no'] for asset in assets], [
            'T/2026/001111',
            'T/2026/002222',
        ])

    def test_xlsx_extracts_only_images_anchored_in_mark_column(self):
        xlsx_path = self.root / 'sample.xlsx'
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Application No.', 'Mark', 'Notes'])
        sheet.append(['T/2026/002509', 'Example mark', 'not a logo'])
        mark_image = SpreadsheetImage(str(self.image_path))
        mark_image.anchor = 'B2'
        sheet.add_image(mark_image)
        unrelated_image = SpreadsheetImage(str(self.image_path))
        unrelated_image.anchor = 'C2'
        sheet.add_image(unrelated_image)
        workbook.save(xlsx_path)
        workbook.close()

        assets = xlsx_assets(xlsx_path, self.root, lambda *_: None)

        self.assertEqual(len(assets), 1)
        self.assertEqual(assets[0]['column'], 2)
        self.assertEqual(assets[0]['mark'], 'Example mark')
        self.assertEqual(assets[0]['field_label'], 'Mark column logo')

    def test_prepared_verification_preserves_existing_match_result(self):
        image = cv(self.image_path)

        self.assertEqual(
            verify(image, image),
            verify_prepared(prepare_verification(image), prepare_verification(image)),
        )

    def test_xlsx_with_images_but_no_mark_header_reports_the_problem(self):
        xlsx_path = self.root / 'unlabeled.xlsx'
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Name', 'Notes'])
        sheet.append(['Example', 'not a logo'])
        image = SpreadsheetImage(str(self.image_path))
        image.anchor = 'B2'
        sheet.add_image(image)
        workbook.save(xlsx_path)
        workbook.close()

        with self.assertRaisesRegex(ValueError, 'has no Mark/logo/device column'):
            xlsx_assets(xlsx_path, self.root, lambda *_: None)

    def test_xlsx_reports_vector_images_that_cannot_be_rendered(self):
        xlsx_path = self.root / 'unrenderable.xlsx'
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Application No.', 'Mark'])
        sheet.append(['T/2026/002509', 'Example mark'])
        workbook.save(xlsx_path)
        workbook.close()

        with (
            patch('app.engine.ooxml_images', return_value=[('Sheet', 2, 2, b'wmf-data', '.wmf')]),
            patch('app.engine.render_vector_bytes', return_value=None),
        ):
            with self.assertRaisesRegex(ValueError, 'Install ImageMagick or Inkscape'):
                xlsx_assets(xlsx_path, self.root, lambda *_: None)

    def test_pdf_to_xlsx_logo_match_works_in_either_upload_order(self):
        pdf_path = self.root / 'government.pdf'
        doc = pymupdf.open()
        page = doc.new_page(width=300, height=300)
        page.insert_text((20, 40), '(540)')
        page.insert_image(pymupdf.Rect(120, 50, 200, 110), filename=str(self.image_path))
        page.insert_text((20, 140), '(550)')
        doc.save(pdf_path)
        doc.close()

        xlsx_path = self.root / 'office.xlsx'
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(['Application No.', 'Mark'])
        sheet.append(['T/2026/002509', 'Example mark'])
        sheet.add_image(SpreadsheetImage(str(self.image_path)), 'B2')
        workbook.save(xlsx_path)
        workbook.close()

        for index, (source_a, source_b, field_a, field_b) in enumerate((
            (pdf_path, xlsx_path, '(540) logo', 'Mark column logo'),
            (xlsx_path, pdf_path, 'Mark column logo', '(540) logo'),
        )):
            job_dir = self.root / f'job-{index}'
            summary = run(job_dir, source_a, source_b, lambda *_: None)
            matches = json.loads((job_dir / 'report' / 'matches.json').read_text())

            self.assertEqual(summary['matched'], 1)
            self.assertEqual(matches[0]['source_a_field'], field_a)
            self.assertEqual(matches[0]['source_b_field'], field_b)


if __name__ == '__main__':
    unittest.main()
