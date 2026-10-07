from pathlib import Path
from functools import lru_cache
import csv, hashlib, json, re, traceback, zipfile, io, math, shutil, subprocess, tempfile, sys
import pymupdf
import cv2
import numpy as np
from PIL import Image
from openpyxl import load_workbook
from skimage.metrics import structural_similarity as ssim
import xml.etree.ElementTree as ET

MEDIA_EXT={'.png','.jpg','.jpeg','.bmp','.webp','.tif','.tiff','.gif','.wmf','.emf','.svg'}
APP_RE=re.compile(r'(?<![A-Z0-9])(?:[A-Z]{1,8}/)?[T7I]\s*/\s*(\d{4})\s*/\s*(\d{3,10})(?!\d)',re.I)
MARK_WORDS={'MARK','LOGO','DEVICE','TRADEMARK','TRADEMARKS','LABEL','BRAND','FIGURE'}
PDF_FIELD_CODE_RE=re.compile(r'\(\d{3}\)')
MARK_HEADERS={word.lower() for word in MARK_WORDS}|{'reproductionofmark'}

def norm_app(v):
    if v is None:return None
    m=APP_RE.search(str(v).upper())
    return f'T/{m.group(1)}/{m.group(2)}' if m else None

def sha_file(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for b in iter(lambda:f.read(1048576),b''):h.update(b)
    return h.hexdigest()

def cv(data):
    try:
        if isinstance(data,(bytes,bytearray)):return cv2.imdecode(np.frombuffer(data,np.uint8),cv2.IMREAD_COLOR)
        return cv2.imdecode(np.fromfile(str(data),np.uint8),cv2.IMREAD_COLOR)
    except Exception:return None

def render_windows_metafile(src,dst):
    if sys.platform!='win32':return None
    import ctypes
    from ctypes import wintypes

    class StartupInput(ctypes.Structure):
        _fields_=[
            ('GdiplusVersion',wintypes.UINT),
            ('DebugEventCallback',ctypes.c_void_p),
            ('SuppressBackgroundThread',wintypes.BOOL),
            ('SuppressExternalCodecs',wintypes.BOOL),
        ]

    gdiplus=ctypes.WinDLL('gdiplus')
    token=ctypes.c_size_t()
    startup=StartupInput(1,None,False,False)
    if gdiplus.GdiplusStartup(ctypes.byref(token),ctypes.byref(startup),None)!=0:return None
    metafile=ctypes.c_void_p()
    bitmap=ctypes.c_void_p()
    graphics=ctypes.c_void_p()
    try:
        if gdiplus.GdipCreateMetafileFromFile(str(src),ctypes.byref(metafile))!=0:return None
        width=wintypes.UINT();height=wintypes.UINT()
        if gdiplus.GdipGetImageWidth(metafile,ctypes.byref(width))!=0:return None
        if gdiplus.GdipGetImageHeight(metafile,ctypes.byref(height))!=0:return None
        if not width.value or not height.value:return None
        if width.value>12000 or height.value>12000:return None
        if gdiplus.GdipCreateBitmapFromScan0(width.value,height.value,0,0x26200A,None,ctypes.byref(bitmap))!=0:return None
        if gdiplus.GdipGetImageGraphicsContext(bitmap,ctypes.byref(graphics))!=0:return None
        if gdiplus.GdipGraphicsClear(graphics,0xFFFFFFFF)!=0:return None
        if gdiplus.GdipDrawImageRectI(graphics,metafile,0,0,width.value,height.value)!=0:return None
        png_encoder=ctypes.c_byte*16
        encoder_id=png_encoder(0x06,0xF4,0x7C,0x55,0x04,0x1A,0xD3,0x11,0x9A,0x73,0x00,0x00,0xF8,0x1E,0xF3,0x2E)
        if gdiplus.GdipSaveImageToFile(bitmap,str(dst),ctypes.byref(encoder_id),None)!=0:return None
        return cv(dst)
    finally:
        if graphics:gdiplus.GdipDeleteGraphics(graphics)
        if bitmap:gdiplus.GdipDisposeImage(bitmap)
        if metafile:gdiplus.GdipDisposeImage(metafile)
        gdiplus.GdiplusShutdown(token)

def render_vector_bytes(data,ext):
    suffix=ext if ext.startswith('.') else '.'+ext
    with tempfile.TemporaryDirectory() as td:
        src=Path(td)/('asset'+suffix);dst=Path(td)/'asset.png';src.write_bytes(data)
        if suffix in {'.emf','.wmf'}:
            im=render_windows_metafile(src,dst)
            if im is not None:return im
        magick=shutil.which('magick');ink=shutil.which('inkscape')
        cmds=[]
        if magick:cmds.append([magick,'-density','300',str(src),'-background','white','-alpha','remove','-alpha','off',str(dst)])
        if ink and suffix in {'.svg','.emf','.wmf'}:cmds.append([ink,str(src),'--export-type=png','--export-filename='+str(dst)])
        for cmd in cmds:
            try:
                p=subprocess.run(cmd,capture_output=True,timeout=45)
                if p.returncode==0 and dst.exists():
                    im=cv(dst)
                    if im is not None:return im
            except Exception:pass
    return None

def gray(img):return cv2.cvtColor(img,cv2.COLOR_BGR2GRAY) if len(img.shape)==3 else img.copy()

def logo_mask(g):
    g=cv2.normalize(g,None,0,255,cv2.NORM_MINMAX);g=cv2.GaussianBlur(g,(3,3),0)
    _,m=cv2.threshold(g,0,255,cv2.THRESH_BINARY_INV+cv2.THRESH_OTSU)
    if np.mean(g)<127:m=cv2.bitwise_not(m)
    m=cv2.morphologyEx(m,cv2.MORPH_OPEN,cv2.getStructuringElement(cv2.MORPH_ELLIPSE,(3,3)))
    return m

def canonical(img,size=256):
    if img is None:return None
    g=gray(img);m=logo_mask(g)
    n,lab,stats,_=cv2.connectedComponentsWithStats(m,8);clean=np.zeros_like(m);min_area=max(6,int(m.size*.00003))
    for i in range(1,n):
        if stats[i,cv2.CC_STAT_AREA]>=min_area:clean[lab==i]=255
    ys,xs=np.where(clean>0)
    if len(xs)>20:
        pad=max(4,int(min(g.shape)*.03));g=g[max(0,ys.min()-pad):min(g.shape[0],ys.max()+pad+1),max(0,xs.min()-pad):min(g.shape[1],xs.max()+pad+1)]
    g=cv2.resize(g,(size,size),interpolation=cv2.INTER_AREA)
    if np.mean(g)<127:g=255-g
    return g

def bits(a):
    v=0
    for x in a.flatten().astype(np.uint8):v=(v<<1)|int(x)
    return f'{v:0{(a.size+3)//4}x}'

def hamming(a,b):return (int(a,16)^int(b,16)).bit_count()

def hashes(g):
    sm=cv2.resize(g,(32,32),interpolation=cv2.INTER_AREA).astype(np.float32);d=cv2.dct(sm)[:8,:8];med=np.median(d.flatten()[1:])
    ph=bits((d>med).astype(np.uint8));dhg=cv2.resize(g,(9,8),interpolation=cv2.INTER_AREA);dh=bits((dhg[:,1:]>dhg[:,:-1]).astype(np.uint8));ahg=cv2.resize(g,(8,8),interpolation=cv2.INTER_AREA);ah=bits((ahg>ahg.mean()).astype(np.uint8))
    return ph,dh,ah

def fingerprint_from_canonical(g):
    m=logo_mask(g)
    ph,dh,ah=hashes(g)
    return {'sha256':hashlib.sha256(g.tobytes()).hexdigest(),'mask_sha256':hashlib.sha256(m.tobytes()).hexdigest(),'phash':ph,'dhash':dh,'ahash':ah}

def fp(img):
    g=canonical(img)
    return fingerprint_from_canonical(g) if g is not None else {}

def fast_score(a,b):
    ga,gb=canonical(a),canonical(b)
    if ga is None or gb is None:return 0
    fa,fb=hashes(ga),hashes(gb)
    return .55*(1-hamming(fa[0],fb[0])/64)+.30*(1-hamming(fa[1],fb[1])/64)+.15*(1-hamming(fa[2],fb[2])/64)

def contour_signature(img):
    g=gray(img);m=logo_mask(g);contours,_=cv2.findContours(m,cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)
    if not contours:return None
    contour=max(contours,key=cv2.contourArea);hu=cv2.HuMoments(cv2.moments(contour)).flatten();hu=np.sign(hu)*np.log10(np.abs(hu)+1e-30);x,y,w,h=cv2.boundingRect(contour)
    return hu,cv2.contourArea(contour)/(g.shape[0]*g.shape[1]),w/g.shape[1],h/g.shape[0]

def contour_similarity_from_signatures(x,y):
    if not x or not y:return 0
    hu=.0
    for i in range(7):hu+=min(abs(float(x[0][i]-y[0][i]))/10,1)
    hu=1-hu/7;ar=1-abs(x[1]-y[1])/max(x[1],y[1],1e-6);sh=1-(abs(x[2]-y[2])+abs(x[3]-y[3]))/2
    return max(0,min(1,.55*hu+.25*ar+.20*sh))

def contour_similarity(a,b):
    return contour_similarity_from_signatures(contour_signature(a),contour_signature(b))

def prepare_verification(img):
    g=canonical(img)
    if g is None:return {}
    sift=cv2.SIFT_create(nfeatures=600);keypoints,descriptors=sift.detectAndCompute(g,None)
    return {
        'canonical':g,
        'fingerprint':fingerprint_from_canonical(g),
        'mask':logo_mask(g),
        'keypoints':keypoints,
        'descriptors':descriptors,
        'contour':contour_signature(img),
    }

def verify_prepared(first,second):
    if not first or not second:return 0,{},False
    ga,gb=first['canonical'],second['canonical']
    fa,fb=first['fingerprint'],second['fingerprint'];ph=1-hamming(fa['phash'],fb['phash'])/64;dh=1-hamming(fa['dhash'],fb['dhash'])/64;ah=1-hamming(fa['ahash'],fb['ahash']);ss=float(ssim(ga,gb,data_range=255));cs=contour_similarity_from_signatures(first['contour'],second['contour'])
    k1,d1=first['keypoints'],first['descriptors'];k2,d2=second['keypoints'],second['descriptors'];good=[];inl=0
    if d1 is not None and d2 is not None:
        for pair in cv2.BFMatcher().knnMatch(d1,d2,k=2):
            if len(pair)==2 and pair[0].distance<.72*pair[1].distance:good.append(pair[0])
        if len(good)>=4:
            src=np.float32([k1[m.queryIdx].pt for m in good]).reshape(-1,1,2);dst=np.float32([k2[m.trainIdx].pt for m in good]).reshape(-1,1,2)
            _,mask=cv2.findHomography(src,dst,cv2.RANSAC,4);inl=int(mask.sum()) if mask is not None else 0
    ma,mb=first['mask'],second['mask'];best_iou=0;best_ms=0
    for ang in (0,-2,2,-4,4):
        M=cv2.getRotationMatrix2D((128,128),ang,1);rb=cv2.warpAffine(mb,M,(256,256),borderValue=0);inter=np.logical_and(ma>0,rb>0).sum();union=np.logical_or(ma>0,rb>0).sum();iou=inter/max(union,1);ms=float(ssim(ma,rb,data_range=255));
        if iou>best_iou or (iou==best_iou and ms>best_ms):best_iou,best_ms=iou,ms
    exact=fa['sha256']==fb['sha256'] or (best_iou>=.92 and best_ms>=.995 and ph>=.94 and ss>=.94)
    final=.18*max(ph,0)+.10*max(dh,0)+.05*max(ah,0)+.20*max(ss,0)+.17*cs+.30*min(1,inl/12)
    if exact:final=max(final,.995)
    return float(final),{'phash':round(ph,4),'dhash':round(dh,4),'ahash':round(ah,4),'ssim':round(ss,5),'contour_similarity':round(cs,5),'mask_iou':round(best_iou,5),'mask_ssim':round(best_ms,5),'sift_good':len(good),'sift_inliers':inl,'canonical_sha256':fa['sha256'],'mask_sha256':fa['mask_sha256']},exact

def verify(a,b):
    return verify_prepared(prepare_verification(a),prepare_verification(b))

def app_metadata(text):
    m=APP_RE.search(text or '')
    return f'T/{m.group(1)}/{m.group(2)}' if m else ''

def page_applications(page):
    applications=[]
    for block in page.get_text('blocks'):
        text=block[4] if len(block)>4 else ''
        for match in APP_RE.finditer(text or ''):
            applications.append(((block[1]+block[3])/2,f'T/{match.group(1)}/{match.group(2)}'))
    return applications

def page_text_words(page):
    try:return page.get_text('words')
    except:return []

def pdf_assets(path,out,progress):
    doc=pymupdf.open(str(path));assets=[];pages=len(doc);asset_no=0
    for pi,page in enumerate(doc):
        progress(int(pi/max(pages,1)*40),f'Extracting PDF: page {pi+1}/{pages}')
        fields=page.search_for('(540)')
        applications=page_applications(page)
        if not fields:
            progress(int((pi+1)/max(pages,1)*40),f'Extracting PDF: page {pi+1}/{pages}')
            continue
        words=page.get_text('words')
        code_words=[word for word in words if PDF_FIELD_CODE_RE.fullmatch(word[4]) and word[0]<120]
        image_rects=[]
        for info in page.get_images(full=True):
            for rect in page.get_image_rects(info[0],transform=False):
                if rect.width<10 or rect.height<10 or rect.width*rect.height>page.rect.width*page.rect.height*.90:continue
                image_rects.append(rect)
        vector_rects=[drawing['rect'] for drawing in page.get_drawings() if drawing.get('rect')]
        for field in fields:
            next_codes=[word[1] for word in code_words if word[1]>=field.y1-2 and word[1]>field.y0]
            field_bottom=min(next_codes,default=page.rect.y1)
            field_top=max(page.rect.y0,field.y0-8)
            candidates=[
                rect for rect in image_rects
                if rect.x0>=field.x1+20 and field_top<=((rect.y0+rect.y1)/2)<field_bottom
            ]
            method='pdf_540_embedded_image'
            if candidates:
                bbox=pymupdf.Rect(
                    min(rect.x0 for rect in candidates),min(rect.y0 for rect in candidates),
                    max(rect.x1 for rect in candidates),max(rect.y1 for rect in candidates),
                )
            else:
                candidates=[
                    rect for rect in vector_rects
                    if rect.width>2 and rect.height>2
                    and rect.width<page.rect.width*.7 and rect.height<page.rect.height*.3
                    and rect.x0>=field.x1+20 and field_top<=((rect.y0+rect.y1)/2)<field_bottom
                ]
                if not candidates:continue
                bbox=pymupdf.Rect(
                    min(rect.x0 for rect in candidates),min(rect.y0 for rect in candidates),
                    max(rect.x1 for rect in candidates),max(rect.y1 for rect in candidates),
                )
                method='pdf_540_vector_art'
            bbox=pymupdf.Rect(
                max(page.rect.x0,bbox.x0-2),max(page.rect.y0,bbox.y0-2),
                min(page.rect.x1,bbox.x1+2),min(page.rect.y1,bbox.y1+2),
            )
            pix=page.get_pixmap(matrix=pymupdf.Matrix(2,2),clip=bbox,alpha=False)
            im=cv(pix.tobytes('png'))
            if im is None:raise ValueError(f'Could not render (540) logo on PDF page {pi+1}')
            asset_no+=1;fn=f'A_{asset_no:06d}.png';cv2.imwrite(str(out/fn),im)
            application_no=(
                min(applications,key=lambda item:abs(item[0]-(field.y0+field.y1)/2))[1]
                if applications else app_metadata(page.get_text('text'))
            )
            assets.append({'asset_id':f'A-{asset_no:06d}','source_type':'pdf','source_file':path.name,'page':pi+1,'bbox':[bbox.x0,bbox.y0,bbox.x1,bbox.y1],'application_no':application_no,'field_code':'(540)','field_label':'(540) logo','image_file':fn,'extract_method':method})
        progress(int((pi+1)/max(pages,1)*40),f'Extracting PDF: page {pi+1}/{pages}')
    return assets

def ooxml_images(path):
    result=[]
    with zipfile.ZipFile(path) as z:
        names=set(z.namelist());ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';nsr='http://schemas.openxmlformats.org/officeDocument/2006/relationships';xdr='http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing';a='http://schemas.openxmlformats.org/drawingml/2006/main'
        wb=ET.fromstring(z.read('xl/workbook.xml'));rels=ET.fromstring(z.read('xl/_rels/workbook.xml.rels'));wr={x.attrib.get('Id'):x.attrib.get('Target','') for x in rels}
        for sh in wb.findall('{%s}sheets/{%s}sheet'%(ns,ns)):
            rid=sh.attrib.get('{%s}id'%nsr);target=wr.get(rid)
            if not target:continue
            sp='xl/'+target.split('xl/')[-1].lstrip('/');sprels='xl/worksheets/_rels/'+Path(sp).name+'.rels'
            if sprels not in names:continue
            sr=ET.fromstring(z.read(sp));rr=ET.fromstring(z.read(sprels));rmap={x.attrib.get('Id'):x.attrib.get('Target','') for x in rr}
            de=sr.find('.//{%s}drawing'%ns)
            if de is None:continue
            dp='xl/drawings/'+Path(rmap.get(de.attrib.get('{%s}id'%nsr),'')).name
            if dp not in names:continue
            dr=ET.fromstring(z.read(dp));drels='xl/drawings/_rels/'+Path(dp).name+'.rels';drr=ET.fromstring(z.read(drels));dmap={x.attrib.get('Id'):x.attrib.get('Target','') for x in drr}
            for anc in dr:
                frm=anc.find('{%s}from'%xdr);blip=anc.find('.//{%s}blip'%a)
                if frm is None or blip is None:continue
                row=int(frm.findtext('{%s}row'%xdr,default='0'))+1;col=int(frm.findtext('{%s}col'%xdr,default='0'))+1;rid2=blip.attrib.get('{%s}embed'%nsr);target2=dmap.get(rid2,'').replace('..','').lstrip('/');media='xl/'+target2.split('xl/')[-1].lstrip('/')
                if media in names:result.append((sh.attrib.get('name','Sheet'),row,col,z.read(media),Path(media).suffix.lower()))
    return result

def xlsx_assets(path,out,progress):
    progress(45,'Reading workbook values and indexing embedded images')
    raw=ooxml_images(path);wb=load_workbook(path,data_only=True,read_only=True,keep_links=False);assets=[];asset_no=0
    image_sheets={image[0] for image in raw}
    try:
        for ws in wb.worksheets:
            max_row=ws.max_row or 0;max_col=max(ws.max_column or 0,1)
            mark_col=None;headers={}
            header_end=min(20,max_row)
            if header_end:
                for r,row_values in enumerate(ws.iter_rows(min_row=1,max_row=header_end,max_col=max_col,values_only=True),1):
                    for c,value in enumerate(row_values,1):
                        v=re.sub(r'[^a-z0-9]+','',str(value or '').lower())
                        if v in MARK_HEADERS and mark_col is None:mark_col=c
                        if 'app' in v and ('no' in v or 'number' in v):headers.setdefault('app',c)
                        if any(k in v for k in ('company','applicant','owner')):headers.setdefault('company',c)
            if mark_col is None:
                if ws.title in image_sheets:raise ValueError(f"Worksheet '{ws.title}' contains images but has no Mark/logo/device column")
                continue
            byrow={}
            for sh,row,col,data,ext in raw:
                if sh!=ws.title or col!=mark_col or row>max_row:continue
                im=cv(data)
                if im is None and ext in {'.wmf','.emf','.svg'}: im=render_vector_bytes(data,ext)
                if im is None:
                    message=f"Could not decode embedded {ext.lstrip('.').upper()} image in worksheet '{ws.title}' at row {row}, column {col}."
                    if ext in {'.wmf','.emf','.svg'}:message+=' Install ImageMagick or Inkscape with support for this format and retry.'
                    else:message+=' The image may be damaged or use an unsupported format.'
                    raise ValueError(message)
                byrow.setdefault(row,[]).append((im,row,col,ext))
            if not byrow:continue
            value_columns=[mark_col,headers.get('app',1)]
            if max_col>=2:value_columns.append(headers.get('company',2))
            values_by_row={}
            first_row=min(byrow);last_row=max(byrow);value_max_col=max(value_columns)
            for row,row_values in enumerate(ws.iter_rows(min_row=first_row,max_row=last_row,max_col=value_max_col,values_only=True),first_row):
                if row in byrow:values_by_row[row]=row_values
            for row,items in byrow.items():
                row_values=values_by_row[row]
                cell=lambda column:row_values[column-1] if column<=len(row_values) else None
                app=str(cell(headers.get('app',1)) or '');company=str(cell(headers.get('company',2)) or '') if max_col>=2 else '';mark=str(cell(mark_col) or '')
                for im,rr,col,ext in items:
                    asset_no+=1;fn=f'B_{asset_no:06d}.png';cv2.imwrite(str(out/fn),im);assets.append({'asset_id':f'B-{asset_no:06d}','source_type':'xlsx','source_file':path.name,'sheet':ws.title,'row':row,'column':col,'application_no':norm_app(app) or app,'company':company,'mark':mark,'field_code':'Mark','field_label':'Mark column logo','image_file':fn,'extract_method':'xlsx_mark_column_image','media_ext':ext})
    finally:
        wb.close()
    progress(50,f'Extracting XLSX: {len(assets)} images')
    return assets

def extract(path,side,out,progress):
    ext=path.suffix.lower()
    return pdf_assets(path,out,progress) if ext=='.pdf' else xlsx_assets(path,out,progress)

def run(job_dir,file_a,file_b,update):
    work=job_dir/'assets';work.mkdir(parents=True,exist_ok=True);report=job_dir/'report';report.mkdir(parents=True,exist_ok=True)
    update(1,'Checking whether source files are identical')
    hash_a=sha_file(file_a);hash_b=sha_file(file_b)
    identical_sources=hash_a==hash_b
    A=extract(file_a,'A',work,update)
    if identical_sources:
        update(45,'Same file: reusing extracted logo images for comparison')
        B=[
            {
                **asset,
                'asset_id':f"B-{index+1:06d}",
                'source_file':file_b.name,
            }
            for index,asset in enumerate(A)
        ]
        matches=[];matched_b=set();total=len(A)
        metrics={'phash':1.0,'dhash':1.0,'ahash':1.0,'ssim':1.0,'contour_similarity':1.0,'mask_iou':1.0,'mask_ssim':1.0,'sift_inliers':0}
        for idx,(a,b) in enumerate(zip(A,B)):
            matched_b.add(b['asset_id'])
            matches.append({
                'match_id':f'M-{len(matches)+1:06d}',
                'source_a_asset_id':a['asset_id'],
                'source_b_asset_id':b['asset_id'],
                'source_a_file':file_a.name,
                'source_b_file':file_b.name,
                'source_a_page':a.get('page'),
                'source_b_page':b.get('page'),
                'source_a_sheet':a.get('sheet'),
                'source_a_row':a.get('row'),
                'source_b_sheet':b.get('sheet'),
                'source_b_row':b.get('row'),
                'source_a_field':a.get('field_label',''),
                'source_b_field':b.get('field_label',''),
                'application_a':a.get('application_no',''),
                'application_b':b.get('application_no',''),
                'application_no_equal':bool(a.get('application_no') and b.get('application_no') and a.get('application_no')==b.get('application_no')),
                'score':1.0,
                'decision':'EXACT_VISUAL_IDENTITY',
                'match_origin':'identical_source_image',
                'metrics':metrics,
                'asset_a_url':f"/api/jobs/{job_dir.name}/assets/A/{a['image_file']}",
                'asset_b_url':f"/api/jobs/{job_dir.name}/assets/B/{b['image_file']}",
            })
            update(45+int((idx+1)/max(total,1)*50),f'Matching logos from identical files: {idx+1}/{total}')
    else:
        B=extract(file_b,'B',work,update)
        update(55,f'Building visual index: {len(B)} assets')
        Bfp=[];bysha={};buckets={}
        for b in B:
            im=cv(work/b['image_file']);f=fp(im);b['_fp']=f;Bfp.append(b)
            if f.get('sha256'):bysha.setdefault(f['sha256'],[]).append(b)
            for key in (f.get('phash'),f.get('dhash'),f.get('ahash')):
                if key:buckets.setdefault(key[:4],[]).append(b)
        @lru_cache(maxsize=128)
        def prepared_file(image_file):
            return prepare_verification(cv(work/image_file))
        matches=[];matched_b=set();total=len(A)
        for idx,a in enumerate(A):
            aim=cv(work/a['image_file']);af=fp(aim);prepared_a=None;cands=[];seen=set()
            for b in bysha.get(af.get('sha256'),[]):cands.append((1.0,b,'exact_hash'));seen.add(b['asset_id'])
            keys=[]
            for key in (af.get('phash'),af.get('dhash'),af.get('ahash')):
                if key:keys.append(key[:4])
            for key in keys:
                for b in buckets.get(key,[]):
                    if b['asset_id'] not in seen:
                        f=b['_fp'];score=.55*(1-hamming(af['phash'],f['phash'])/64)+.30*(1-hamming(af['dhash'],f['dhash'])/64)+.15*(1-hamming(af['ahash'],f['ahash'])/64)
                        cands.append((score,b,'hash_candidate'));seen.add(b['asset_id'])
            # Avoid an exhaustive quadratic scan for large catalogs.
            if len(cands)<12 and len(Bfp)<=1000:
                top=[]
                for b in Bfp:
                    if b['asset_id'] in seen:continue
                    f=b['_fp'];s=.55*(1-hamming(af['phash'],f['phash'])/64)+.30*(1-hamming(af['dhash'],f['dhash'])/64)+.15*(1-hamming(af['ahash'],f['ahash'])/64)
                    if s>=.58:top.append((s,b,'hash_candidate'))
                top.sort(key=lambda x:x[0],reverse=True);cands.extend(top[:30])
            cands.sort(key=lambda x:x[0],reverse=True)
            exact_candidates=[candidate for candidate in cands if candidate[2]=='exact_hash']
            review_candidates=exact_candidates+[candidate for candidate in cands if candidate[2]!='exact_hash'][:10]
            for _,b,origin in review_candidates:
                if origin=='exact_hash':
                    score=1.0;exact=True
                    detail={'phash':1.0,'dhash':1.0,'ahash':1.0,'ssim':1.0,'contour_similarity':1.0,'mask_iou':1.0,'mask_ssim':1.0,'sift_good':0,'sift_inliers':0,'canonical_sha256':af['sha256'],'mask_sha256':af['mask_sha256']}
                else:
                    if prepared_a is None:prepared_a=prepare_verification(aim)
                    score,detail,exact=verify_prepared(prepared_a,prepared_file(b['image_file']))
                if exact:decision='EXACT_VISUAL_IDENTITY'
                elif score>=.82:decision='VERY_HIGH_VISUAL_SIMILARITY'
                elif score>=.70:decision='HIGH_VISUAL_SIMILARITY'
                else:continue
                matched_b.add(b['asset_id'])
                matches.append({'match_id':f'M-{len(matches)+1:06d}','source_a_asset_id':a['asset_id'],'source_b_asset_id':b['asset_id'],'source_a_file':a['source_file'],'source_b_file':b['source_file'],'source_a_page':a.get('page'),'source_b_page':b.get('page'),'source_a_sheet':a.get('sheet'),'source_a_row':a.get('row'),'source_b_sheet':b.get('sheet'),'source_b_row':b.get('row'),'source_a_field':a.get('field_label',''),'source_b_field':b.get('field_label',''),'application_a':a.get('application_no',''),'application_b':b.get('application_no',''),'application_no_equal':bool(a.get('application_no') and b.get('application_no') and a.get('application_no')==b.get('application_no')),'score':round(score,6),'decision':decision,'match_origin':origin,'metrics':detail,'asset_a_url':f"/api/jobs/{job_dir.name}/assets/A/{a['image_file']}",'asset_b_url':f"/api/jobs/{job_dir.name}/assets/B/{b['image_file']}"})
            update(55+int((idx+1)/max(total,1)*40),f'Matching logos: {idx+1}/{total}')
    for a in A:a.pop('_fp',None)
    for b in B:b.pop('_fp',None)
    note=(
        'Identical source files: logo images were extracted once and matched directly.'
        if identical_sources
        else 'Application numbers are metadata only. Visual matching is independent of application-number equality. Similarity results are technical evidence for review, not a legal infringement verdict.'
    )
    summary={'source_a':{'file':file_a.name,'type':file_a.suffix.lower(),'assets':len(A),'sha256':hash_a},'source_b':{'file':file_b.name,'type':file_b.suffix.lower(),'assets':len(B),'sha256':hash_b},'matched':len(matches),'unmatched_a':len(A)-len({m['source_a_asset_id'] for m in matches}),'unmatched_b':len(B)-len(matched_b),'exact_identity':sum(m['decision']=='EXACT_VISUAL_IDENTITY' for m in matches),'very_high':sum(m['decision']=='VERY_HIGH_VISUAL_SIMILARITY' for m in matches),'high':sum(m['decision']=='HIGH_VISUAL_SIMILARITY' for m in matches),'application_number_equal_matches':sum(m['application_no_equal'] for m in matches),'application_number_used_for_matching':False,'identical_sources':identical_sources,'analysis_skipped':False,'same_file_logo_scan':identical_sources,'engine':'Apexive AI Trademark Visual Conflict Engine v4','message':note,'note':note}
    (report/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf8');(report/'matches.json').write_text(json.dumps(matches,ensure_ascii=False,indent=2),encoding='utf8');(report/'assets_a.json').write_text(json.dumps(A,ensure_ascii=False,indent=2),encoding='utf8');(report/'assets_b.json').write_text(json.dumps(B,ensure_ascii=False,indent=2),encoding='utf8')
    with open(report/'matches.csv','w',newline='',encoding='utf-8-sig') as f:
        fields=['match_id','source_a_asset_id','source_b_asset_id','source_a_file','source_b_file','source_a_page','source_a_sheet','source_a_row','source_b_page','source_b_sheet','source_b_row','source_a_field','source_b_field','application_a','application_b','application_no_equal','score','decision','match_origin'];w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows([{k:m.get(k,'') for k in fields} for m in matches])
    update(100,'Completed')
    return summary
